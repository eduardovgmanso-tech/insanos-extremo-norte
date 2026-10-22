"""
Painel Insanos na nuvem (GitHub Actions).
1) Lê a planilha do Google (conta de serviço) e monta uma planilha temporária só com as abas que o painel usa.
2) Roda o mesmo programa do painel (painel.py) e publica neste repositório.
Nada de dados pessoais fica no repositório: a planilha temporária existe só durante a execução.
"""
import os, sys, json, datetime, unicodedata

PASTA = os.path.dirname(os.path.abspath(__file__))
ABAS = ["Presenca", "Insights", "Dados Integrantes", "PGTO", "Pagamentos"]
COLS_DATA = {"Presenca": [1], "Insights": [1], "Dados Integrantes": [9, 17]}   # colunas de data (0 = A): B, J e R

def _serial_para_data(v):
    try:
        return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))
    except Exception:
        return v

def montar_planilha(valores_por_aba, destino):
    """valores_por_aba: {aba: [[células...]]} (valores já calculados) -> .xlsx temporário para o painel ler."""
    import openpyxl
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for aba, linhas in valores_por_aba.items():
        ws = wb.create_sheet(aba)
        cd = COLS_DATA.get(aba, [])
        for r, linha in enumerate(linhas, start=1):
            for c, v in enumerate(linha, start=1):
                if v == "" or v is None:
                    continue
                if (c - 1) in cd and isinstance(v, (int, float)) and r >= (2 if aba == "Dados Integrantes" else 8) and 20000 < v < 80000:
                    v = _serial_para_data(v)
                ws.cell(r, c).value = v
    wb.save(destino)

def api_google(cred_json):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    cred = service_account.Credentials.from_service_account_info(json.loads(cred_json),
           scopes=["https://www.googleapis.com/auth/spreadsheets"])
    return build("sheets", "v4", credentials=cred, cache_discovery=False)

def ler_google(planilha_id, cred_json, api=None):
    api = api or api_google(cred_json)
    meta = api.spreadsheets().get(spreadsheetId=planilha_id, fields="properties.title,sheets.properties.title").execute()
    existentes = [s["properties"]["title"] for s in meta["sheets"]]
    print(f"Planilha do Google: \"{meta.get('properties', {}).get('title', '?')}\" com {len(existentes)} abas: " + ", ".join(f'"{x}"' for x in existentes))
    # acha cada aba sem ligar para espaços, acentos ou maiúsculas ("Presença " = "Presenca")
    norm = lambda x: "".join(ch for ch in unicodedata.normalize("NFD", str(x).strip().lower()) if not unicodedata.combining(ch))
    real = {}
    for a in ABAS:
        achou = next((e for e in existentes if norm(e) == norm(a)), None)
        if achou:
            real[a] = achou
    faltam = [a for a in ABAS if a not in real]
    if faltam:
        print("Aviso: abas não encontradas na planilha do Google:", ", ".join(faltam))
    if "Presenca" not in real or "Dados Integrantes" not in real:
        print("[ERRO] A planilha precisa ter as abas Presenca e Dados Integrantes. Confira se o PLANILHA_ID é o da planilha certa.")
        sys.exit(1)
    nomes = list(real)
    res = api.spreadsheets().values().batchGet(spreadsheetId=planilha_id, ranges=[f"'{real[a]}'!A1:BZ500" for a in nomes],
          valueRenderOption="UNFORMATTED_VALUE", dateTimeRenderOption="SERIAL_NUMBER").execute()
    return {a: vr.get("values", []) for a, vr in zip(nomes, res.get("valueRanges", []))}

# ---------------------------------------------------------------------
#  Fila: lançamentos que o painel mandou direto para a planilha (aba "Fila")
#  Colunas: A data/hora · B quem · C tipo · D dados (JSON) · E situação (pendente/feito/erro) · F resultado
# ---------------------------------------------------------------------
def _sa(x):
    return "".join(ch for ch in unicodedata.normalize("NFD", str(x or "").strip().lower()) if not unicodedata.combining(ch))

def _col(n):   # 1 -> A
    s = ""
    while n:
        n, r = divmod(n - 1, 26); s = chr(65 + r) + s
    return s

def _aba_real(api, pid, nome):
    meta = api.spreadsheets().get(spreadsheetId=pid, fields="sheets.properties.title").execute()
    return next((s["properties"]["title"] for s in meta["sheets"] if _sa(s["properties"]["title"]) == _sa(nome)), None)

def _ler(api, pid, rng):
    return api.spreadsheets().values().get(spreadsheetId=pid, range=rng, valueRenderOption="UNFORMATTED_VALUE",
                                           dateTimeRenderOption="SERIAL_NUMBER").execute().get("values", [])

def _gravar(api, pid, dados):
    if dados:
        api.spreadsheets().values().batchUpdate(spreadsheetId=pid, body={"valueInputOption": "USER_ENTERED",
            "data": [{"range": r, "values": [[v]]} for r, v in dados]}).execute()

def aplicar_evento(api, pid, d):
    """Lançar evento / Insights / Salvar no Excel: data + nome + marcações na aba (Presenca ou Insights)."""
    aba = _aba_real(api, pid, d.get("aba") or "Presenca")
    if not aba:
        return f"aba {d.get('aba')} não encontrada"
    v = _ler(api, pid, f"'{aba}'!A1:BZ60")
    lin = lambda r: v[r - 1] if len(v) >= r else []
    cols, c = {}, 4
    while True:
        nome = str((lin(6) + [""] * 80)[c - 1]).strip()
        if not nome or nome.lower().startswith("presen"):
            break
        cols[_sa(nome)] = c; c += 1
    alvo = datetime.datetime.strptime(d["data"], "%Y-%m-%d").date()
    def data_da_linha(r):
        x = (lin(r) + ["", "", ""])[1]
        if isinstance(x, (int, float)) and x > 20000:
            return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(x))).date()
        return None
    linha = None
    if d.get("atualizar", True):
        linha = next((r for r in range(8, 43) if data_da_linha(r) == alvo and _sa((lin(r) + ["", "", ""])[2]) == _sa(d.get("evento", ""))), None)
    if linha is None:
        linha = next((r for r in range(8, 43) if not str((lin(r) + ["", ""])[1]).strip()), None)
    if linha is None:
        return "aba cheia (linhas 8 a 42)"
    gr = [(f"'{aba}'!B{linha}", alvo.strftime("%d/%m/%Y")), (f"'{aba}'!C{linha}", d.get("evento", ""))]
    conv = {"sim": "SIM", "nao": "NÃO", "não": "NÃO", "staff": "STAFF", "-": "-", "": ""}
    ignorados = []
    for nome, marca in (d.get("marcacoes") or {}).items():
        cc = cols.get(_sa(nome))
        if not cc:
            ignorados.append(nome); continue
        gr.append((f"'{aba}'!{_col(cc)}{linha}", conv.get(_sa(marca), str(marca))))
    _gravar(api, pid, gr)
    return f"gravado na aba {aba}, linha {linha}" + (f" (sem coluna: {', '.join(ignorados)})" if ignorados else "")

def aplicar_pagamento(api, pid, d):
    """Compensação marcada no painel: PGTO coluna C (nunca mexe em quem já está Pago em Dia)."""
    aba = _aba_real(api, pid, "PGTO")
    if not aba:
        return "aba PGTO não encontrada"
    v = _ler(api, pid, f"'{aba}'!A1:C400")
    gr, feitos = [], []
    for nome, st in (d.get("status") or {}).items():
        for i, row in enumerate(v[1:], start=2):
            rot = str((row + ["", ""])[1]); atual = str((row + ["", "", ""])[2])
            if _sa(rot).endswith(_sa(nome)):
                if "pago" in _sa(atual) and "aguard" not in _sa(atual):
                    break
                gr.append((f"'{aba}'!C{i}", st or "")); feitos.append(nome); break
    _gravar(api, pid, gr)
    return "PGTO: " + (", ".join(feitos) if feitos else "nada alterado")

def processar_fila(api, pid):
    aba = _aba_real(api, pid, "Fila")
    if not aba:
        return 0
    v = _ler(api, pid, f"'{aba}'!A1:F2000")
    n = 0
    for i, row in enumerate(v[1:], start=2):
        row = row + [""] * 6
        if _sa(row[4]) != "pendente":
            continue
        try:
            d = json.loads(row[3]); tp = d.get("tipo")
            if tp == "lancamento_presenca":
                res, sit = aplicar_evento(api, pid, d), "feito"
            elif tp == "pagamento_manual":
                res, sit = aplicar_pagamento(api, pid, d), "feito"
            elif tp == "integrante":
                res, sit = "cadastro ainda é aplicado pelo computador (próxima etapa da nuvem)", "aguardando"
            else:
                res, sit = f"tipo desconhecido: {tp}", "erro"
        except Exception as e:
            res, sit = f"{type(e).__name__}: {e}", "erro"
        _gravar(api, pid, [(f"'{aba}'!E{i}", sit), (f"'{aba}'!F{i}", f"{datetime.datetime.now():%d/%m/%Y %H:%M} · {res}")])
        print(f"Fila linha {i}: {sit} · {res}"); n += 1
    return n

# ---------------------------------------------------------------------
#  Memória entre execuções (na nuvem cada execução começa do zero):
#  os arquivos de estado do painel ficam guardados na aba escondida "_estado" da planilha
# ---------------------------------------------------------------------
ESTADO = ["pagamentos_verificacao.json", "pagamentos_manual.json", "Lancamentos/notificacoes_enviadas.json"]

def _aba_estado(api, pid):
    aba = _aba_real(api, pid, "_estado")
    if not aba:
        api.spreadsheets().batchUpdate(spreadsheetId=pid, body={"requests": [
            {"addSheet": {"properties": {"title": "_estado", "hidden": True}}}]}).execute()
        aba = "_estado"
    return aba

def estado_carregar(api, pid):
    aba = _aba_estado(api, pid)
    for row in _ler(api, pid, f"'{aba}'!A1:B20"):
        row = list(row) + ["", ""]
        if row[0] in ESTADO and str(row[1]).strip():
            destino = os.path.join(PASTA, row[0]); os.makedirs(os.path.dirname(destino), exist_ok=True)
            with open(destino, "w", encoding="utf-8") as f:
                f.write(str(row[1]))
    return aba

def estado_salvar(api, pid, aba):
    linhas = []
    for nome in ESTADO:
        caminho = os.path.join(PASTA, nome)
        if os.path.exists(caminho):
            with open(caminho, encoding="utf-8") as f:
                linhas.append([nome, f.read()[:49000]])
    if linhas:
        api.spreadsheets().values().update(spreadsheetId=pid, range=f"'{aba}'!A1:B{len(linhas)}",
            valueInputOption="RAW", body={"values": linhas}).execute()

# ---------------------------------------------------------------------
#  Verificação dos boletos da Iugu (mesma regra do PC), gravando na aba PGTO da planilha do Google
# ---------------------------------------------------------------------
def verificar_boletos(api, pid, g):
    hoje = datetime.date.today()
    ciclo = g.vencimento(hoje.year, hoje.month)
    if hoje > ciclo:
        print(f"Verificação de pagamentos: fora do período (vencimento foi {ciclo:%d/%m}). Nada a fazer.")
        return False
    aba = _aba_real(api, pid, "PGTO")
    if not aba:
        print("Verificação de pagamentos: aba PGTO não encontrada."); return False
    print(f"Verificando pagamentos (ciclo com vencimento em {ciclo:%d/%m/%Y})...")
    v = _ler(api, pid, f"'{aba}'!A1:J400")
    manual = g._manual_carregar(ciclo)
    inativos = g.inativos_planilha()
    itens, grav, novos_pagos = [], [], 0
    for r, row in enumerate(v[1:], start=2):
        row = list(row) + [""] * 10
        link = str(row[9]).strip()
        if not link.startswith("http"):
            continue
        rotulo, atual = str(row[1]).strip(), str(row[2]).strip()
        if any(rotulo.lower().endswith(ap) for ap in inativos):
            print(f"  {rotulo}: afastado/excluído (não verificado)"); continue
        item = {"rotulo": rotulo, "selo": None, "venc_boleto": None, "status": atual, "obs": ""}
        try:
            selo, vb = g._ler_boleto(link)
        except Exception as e:
            item["obs"] = "não consegui abrir o boleto"; itens.append(item); print(f"  {rotulo}: não abriu ({e})"); continue
        item["selo"], item["venc_boleto"] = selo, (vb.isoformat() if vb else None)
        boleto_do_mes = not vb or (vb.year, vb.month) == (ciclo.year, ciclo.month)
        if boleto_do_mes and selo and "pag" in selo.lower():
            novo = "Pago em Dia"
        else:
            novo = "Em aberto"
            if not boleto_do_mes:
                item["obs"] = f"boleto com vencimento {vb:%d/%m/%Y}: o boleto deste mês ainda não foi cadastrado"
            elif selo:
                item["obs"] = f"selo da Iugu: {selo}"
        ap_man = g._manual_de(rotulo, manual)
        if not ap_man and novo != "Pago em Dia" and atual.lower().startswith("aguard"):
            ap_man = rotulo.split()[-1] if rotulo else rotulo
            manual[ap_man] = "Aguardando Compensação"
        if novo == "Pago em Dia":
            if ap_man:
                manual.pop(ap_man, None)
            item["boleto_pago"] = True
        elif ap_man:
            novo = manual[ap_man]
            item["obs"] = (item.get("obs") + " · " if item.get("obs") else "") + "marcado à mão"
        item["status"] = novo
        if novo != atual:
            grav.append((f"'{aba}'!C{r}", novo)); print(f"  {rotulo}: {atual or '(vazio)'} -> {novo}")
            if novo == "Pago em Dia":
                novos_pagos += 1; item["novo_pago"] = True
        itens.append(item)
    _gravar(api, pid, grav)
    g._manual_salvar(ciclo, manual)
    quando = datetime.datetime.now().strftime("%d/%m/%Y às %H:%M")
    ultimo, pagos_em = None, {}
    try:
        with open(g.ARQ_VERIF, encoding="utf-8") as f:
            ant = json.load(f)
        ultimo = ant.get("ultimo_pagamento")
        if str(ant.get("ciclo") or "")[:7] == ciclo.isoformat()[:7]:
            pagos_em = dict(ant.get("pagos_em") or {})
    except Exception:
        pass
    agora_ts = datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
    for it in itens:
        rot = it.get("rotulo", "")
        if it.get("status") == "Pago em Dia":
            pagos_em.setdefault(rot, agora_ts if it.get("novo_pago") else "")
        else:
            pagos_em.pop(rot, None)
    if novos_pagos:
        ultimo = {"ts": agora_ts, "hora": datetime.datetime.now().strftime("%d/%m às %H:%M")}
    with open(g.ARQ_VERIF, "w", encoding="utf-8") as f:
        json.dump({"quando": quando, "ciclo": ciclo.isoformat(), "novos_pagos": novos_pagos, "ultimo_pagamento": ultimo,
                   "pagos_em": pagos_em, "itens": itens}, f, ensure_ascii=False, indent=1)
    print(f"Verificação concluída: {len(itens)} boletos lidos, {len(grav)} status alterado(s).")
    novos = [str(it.get("rotulo", "")).split(" - ")[-1].strip() for it in itens if it.get("novo_pago")]
    if novos:
        hora = datetime.datetime.now().strftime("%d/%m às %H:%M")
        try:
            if len(novos) == 1:
                g.notificar("💰 Pagamento novo", f"{novos[0]} pagou a mensalidade (identificado {hora}).", "pagamento")
            else:
                g.notificar(f"💰 {len(novos)} pagamentos novos", ", ".join(novos) + f" (identificados {hora}).", "pagamento")
        except Exception as e:
            print(f"Aviso: notificação não enviada ({e}).")
    return bool(grav)

def preparar_ambiente():
    """Chaves do cofre -> arquivos temporários que o painel espera."""
    lanc = os.path.join(PASTA, "Lancamentos"); os.makedirs(lanc, exist_ok=True)
    pem, pub = os.environ.get("PUSH_VAPID_PEM", "").strip(), os.environ.get("PUSH_VAPID_PUB", "").strip()
    if pem and pub:
        with open(os.path.join(lanc, "push_vapid_privada.pem"), "w", encoding="utf-8") as f:
            f.write(pem + "\n")
        with open(os.path.join(lanc, "push_vapid.json"), "w", encoding="utf-8") as f:
            f.write(pub)
    # publicação: a chave automática do GitHub, neste mesmo repositório
    dono, repo = (os.environ.get("GITHUB_REPOSITORY", "/") + "/").split("/")[:2]
    tok = os.path.join(PASTA, ".token_github.json")
    with open(tok, "w", encoding="utf-8") as f:
        json.dump({"token": os.environ.get("GITHUB_TOKEN", ""), "usuario": dono, "repo": repo}, f)
    return tok, dono, repo

def main():
    planilha_id, cred = os.environ.get("PLANILHA_ID", "").strip(), os.environ.get("GOOGLE_CREDENTIALS", "")
    if not planilha_id or not cred:
        print("[ERRO] Faltam os segredos PLANILHA_ID e/ou GOOGLE_CREDENTIALS no GitHub."); sys.exit(1)
    evento = os.environ.get("MODO", "")   # schedule (de hora em hora) · repository_dispatch (lançamento) · workflow_dispatch (botão)
    api = api_google(cred)
    aba_estado = estado_carregar(api, planilha_id)
    feitos = processar_fila(api, planilha_id)
    if feitos:
        print(f"Fila: {feitos} lançamento(s) aplicado(s) na planilha.")
    tok, dono, repo = preparar_ambiente()
    xlsx = os.path.join(PASTA, "Insanos.xlsx")
    print("Lendo a planilha do Google...")
    montar_planilha(ler_google(planilha_id, cred, api), xlsx)
    sys.argv = [sys.argv[0], "--publicar"]
    import painel as g
    g.PLANILHA = xlsx
    g.ARQ_TOKEN = tok
    g.GITHUB_USUARIO, g.GITHUB_REPO = dono, repo
    g.SITE_URL = f"https://{dono}.github.io/{repo}/"
    g.limpar_historico_github = lambda: True       # na nuvem o histórico fica (sem limpeza forçada)
    if os.environ.get("LANC_URL") and os.environ.get("LANC_CHAVE"):
        g.NUVEM = {"url": os.environ["LANC_URL"], "chave": os.environ["LANC_CHAVE"]}   # o painel manda os lançamentos direto
    # boletos: nas rodadas de hora em hora (8h às 17h) e no botão "Run workflow"; nos lançamentos, não (para ser rápido)
    hora = datetime.datetime.now().hour
    if evento == "workflow_dispatch" or (evento == "schedule" and 8 <= hora <= 17):
        try:
            if verificar_boletos(api, planilha_id, g):
                print("Relendo a planilha com os pagamentos atualizados...")
                montar_planilha(ler_google(planilha_id, cred, api), xlsx)
        except Exception as e:
            print(f"Aviso: verificação de pagamentos falhou ({type(e).__name__}: {e}). O painel é gerado mesmo assim.")
    try:
        g.main()
    finally:
        try:
            estado_salvar(api, planilha_id, aba_estado)
        except Exception as e:
            print(f"Aviso: não consegui guardar a memória na aba _estado ({e}).")

if __name__ == "__main__":
    main()
